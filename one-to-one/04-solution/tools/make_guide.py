"""Builds the single, self-contained finalisation guide (Markdown + Word) from the verified documents.

    python3 04-solution/tools/make_guide.py        (from the one-to-one folder or anywhere)

Tests, procedures, troubleshooting, the welcome email and accepted risks are extracted from their source files, so the
guide can't drift from them. Output: one-to-one/FINALISATION_GUIDE.md and one-to-one/ESG-1to1-Finalisation-Guide.docx
"""
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))   # one-to-one/


def read(rel):
    with open(os.path.join(ROOT, rel), encoding='utf-8') as f:
        return f.read()


def section(text, start, end=None):
    """Lines from the heading/line starting with `start` up to (not including) the next line starting with `end`."""
    lines, out, on = text.split('\n'), [], False
    for ln in lines:
        if not on and ln.startswith(start):
            on = True
            continue
        if on and end is not None and ln.startswith(end):
            break
        if on:
            out.append(ln)
    if not out:
        raise SystemExit('section not found: ' + start)
    return '\n'.join(out).strip('\n')


def delink(md):
    """Self-contained: keep the text of internal links, drop the target."""
    return re.sub(r'\[([^\]]+)\]\((?!https?:)[^)]+\)', r'\1', md)


def add_result_columns(table_md):
    out = []
    for ln in table_md.split('\n'):
        if ln.startswith('|---'):
            out.append(ln + '---|---|---|')
        elif ln.startswith('| #'):
            out.append(ln + ' Result | Date | Initials |')
        elif ln.startswith('| '):
            out.append(ln + '  |  |  |')
        else:
            out.append(ln)
    return '\n'.join(out)


def demote(md, by=1):
    return re.sub(r'^(#+) ', lambda m: '#' * (len(m.group(1)) + by) + ' ', md, flags=re.M)


plan = read('IMPLEMENTATION_PLAN.md')
H_intro = section(plan, '## H. Testing', '### Functional tests')
F_table = section(plan, '### Functional tests', '### Privacy and bypass tests')
P_table = section(plan, '### Privacy and bypass tests', '---')
risks = '\n'.join(l for l in section(plan, '### J.3 Residual risks', '---').split('\n') if l.startswith('|'))
cfg = read('CONFIG_REFERENCE.md')
names = section(cfg, '## 2. Names you must never change', '## 3.')
procedures = section(cfg, '## 4. Operating procedures', '## 5.')
cfg_trouble = section(cfg, '## 6. Troubleshooting', '## 7.')
imp = read('04-solution/IMPORT_GUIDE.md')
imp_trouble = section(imp, '## 7. Troubleshooting', '## 8.')
legacy = section(imp, '## 6. Fallback: legacy packages', '## 7.')
setup_does = section(imp, '**What setup does:**', '**What setup does not do.**')
welcome = section(read('03-flows/WELCOME_EMAIL.md'), '---')
fix_prompt = section(read('NEXT_STEPS.md'), '## If something fails')
fix_prompt = fix_prompt.split('```')[0] + '```' + fix_prompt.split('```')[1] + '```'   # the prompt only

GUIDE = r'''# ESG Monthly Technician 1-to-1 — Finalisation Guide

**From the end of Stage 0B to a live system, then running it month to month.** This document is self-contained: you don't need any other file to follow it. Tick each box as you go.

Version: system v5.1.0 (frontend) · solution 1.5.0.1 · guide dated 19-Oct-2026

---

## 1. Where you are, and what "complete" means

**Already done:**

- **Stage 0A/0B passed** in the live tenant:
  - the private site exists, with the `OneToOne` page and the Welcome web part;
  - the `1-to-1 Technicians` group exists;
  - the `1-to-1 Edit` and `1-to-1 Read` permission levels exist.
- **Everything to install is built and verified offline:**
  - frontend: 319 automated checks, 108 more for the action tracker, and 45 for the one-time reset tool;
  - flows: simulated end to end, 298 checks, and 42 deliberately broken copies all caught by the validator;
  - independent review rounds fixed 55 defects (see VERIFICATION.md).

**New in 5.1.0 / solution 1.5 – action tracking** (already on an earlier version? §4.2 has the upgrade order):

- **Every agreed action has its own record** (*1-to-1 Actions*), added automatically when you release a review and followed month to month. The tracker starts empty: actions agreed before the upgrade stay in their reviews.
- **Technicians: My actions.** Each person sees only their own actions, sends a quick update at any time, says when an action is waiting on something, and can ask you to close their own actions.
- **You: the Actions page.** What needs you first (close-out requests, overdue, due soon, waiting, no recent progress, problems), filters, one **Update…** dialog (progress, status, deadline with a reason – the original is kept – owner, waiting), agree or decline close-out, close any action, reopen with a reason, **+ New action**.
- **Outstanding actions** on every monthly review, with what changed and optional follow-up outcomes.
- **Emails:** E19 when someone asks you to close an action; E20 five working days before a deadline (to them, copied to you). No overdue emails.
- **Full history** of every action. Nothing is ever overwritten: a save made on out-of-date information is refused and reloaded.

**Also in 5.0.0 / solution 1.4** (if you are coming from 4.5.0):

- **Appraisal periods follow real completion dates.** ⚙ Settings → Team holds, per person, the next appraisal, the next wellbeing date and the appraisal period start; **Mark appraisal completed…** starts a new period of any length. Moving the next appraisal date never moves a review.
- **Missed months don't count** in averages, the forecast, the trend or on-time ("N rated reviews of M months · K missed").
- **Amendments.** A released review is read-only until you click **Amend…**; every change needs a reason, is kept in a restricted audit only you can see, and the person gets a short email (no content).
- **Reminders** for wellbeing (E12) and appraisals (E17) come from the Settings → Team dates: 7 days before, on the day, then weekly while overdue.
- **Expectations:** no limit of 10; Shared always first; they carry on into later years until retired; retired ones can be deleted permanently.
- **Challenges have been removed** (their records are kept, Owners only).
- A **one-time test-data reset tool** for clearing test records before go-live (§7).

**Not yet proven:** nothing has run in your tenant yet.

> **The system is fully implemented when Phase F (go-live) is done and every privacy test (P1–P36) in Phase C has passed.**

| Phase | What | Who's involved | Time |
|---|---|---|---|
| 0 | Get and check the files | you | 5 min |
| A | Import, set up, see the frontend yourself | you only | ≈ 1 h |
| B | Walk through it as a technician | you + Test A + Test B | 1–2 h |
| C | Security and functional test record | you + test accounts (+ 2 IT questions) | 6–8 h |
| D | Clean down | you | 30 min |
| E | Prepare and invite the real team | you | 1–2 h |
| F | Go live | you | 15 min |
| G | First month, then routine | you | ongoing |

---

## 2. What you need

**The three system files** (nothing else goes into SharePoint or Power Automate):

| File | Goes to |
|---|---|
| `ESGOneToOne_1_5_0_1_unmanaged.zip` | Power Automate → **Solutions → Import solution**. Contains the eleven flows (0 Setup, 1–6 and 8–11); the setup flow creates the lists. |
| `one-to-one.html` | SharePoint → *Welcome Web Part - Custom Layouts* library |
| `one-to-one.js` | Same library, alongside the `.html` |

**Desktop only.** The page is built for laptops and desktops; phones are out of scope.

**Only for the one-time test-data reset** (§7, with its own approval): `one-to-one-reset.html` and `one-to-one-reset.js`, uploaded temporarily to the same library and removed afterwards.

**Fallback only:** if your environment can't import solutions (no Dataverse), use `legacy-package-builder.html` instead (§4.8).

**Accounts:**

- **you**, the Team Leader: site owner **and Site Collection Administrator**;
- **Test A** and **Test B**: two test technician accounts with mailboxes.

## Phase 0 — Get and check the files (5 min)

- [ ] **Right zip.** Open the zip, then open `solution.xml` in Notepad. It must show `<Version>1.5.0.1</Version>` and contain `<RootComponent type="380" schemaName="esg_SiteUrl"`. If it doesn't, you have an older download: get the latest one.
- [ ] **Right frontend.** You have `one-to-one.html` and `one-to-one.js` (v5.1.0). After installing, the top menu bar must show **v5.1.0** (right-hand end).
- [ ] **Unchanged files.** Transfer them the usual approved way. Don't open and re-save them in Word.

---

## 4. Phase A — Import, set up and see the frontend yourself (≈ 1 h, nobody else involved)

### 4.1 Tidy up Stage 0

- [ ] **Site contents:** delete the `ZZProbe…` test lists.
- [ ] **Welcome Web Part - Custom Layouts:** delete `probe-0a*` and `probe-0b*`.
- [ ] **Keep** the site, the `OneToOne` page, the `1-to-1 Technicians` group and the two permission levels.
- [ ] **Check:**
  - The site's regional settings are **English (United Kingdom)** and time zone **(UTC) Dublin, Edinburgh, Lisbon, London**. They're under Site information → View all site settings → Regional settings.
  - You are **Site Collection Administrator**.

### 4.2 Import the solution (10 min)

> **Imported 1.5.0.0 and *1 Daily scheduler* wouldn't turn on?** Import `ESGOneToOne_1_5_0_1_unmanaged.zip` over it the same way, then turn Flow 1 on (`04-solution/REPAIR_1_5_0_1.md`). Nothing else changes.
>
> **Already on 1.4.0.0 / page 5.0.0?** Upgrade to 1.5.0.1 in this order, each step with approval: back up Manager Reviews and 1-to-1 Settings; turn off *1 Daily scheduler*; import `ESGOneToOne_1_5_0_1_unmanaged.zip` as an upgrade (flows 9, 10 and 11 arrive Off); run **0 Setup** and wait for *complete*; turn on flows **2, 3, 4, 5, 6, 8, 9, 10 and 11**; upload page 5.1.0 (Replace, Reset layout cache, Republish – the top bar shows **v5.1.0**); turn on Flow 1. Nothing is imported: the tracker starts empty. (Full order, going back and checks T1–T16: `MIGRATION_1_5.md`.)
>
> **Already imported an earlier version (1.3.0.x / page 4.5.0 or before)?** Upgrade in this order (written for 1.4 – use the 1.5 zip and page 5.1.0, and in step 5 also turn on flows 9, 10 and 11), each step with approval:
>
> 1. **Back up:** export each list to CSV (Technicians, Appraisal Expectations, Monthly Reviews, Manager Reviews, Wellbeing Assessments, 1-to-1 Feedback, 1-to-1 Challenges). Keep the old zip and page files.
> 2. **Turn off** *7 Challenge* (if you have it) and *1 Daily scheduler*.
> 3. Import `ESGOneToOne_1_5_0_1_unmanaged.zip` over it the same way: it says it will **upgrade**; leave *"Enable Plugin steps and flows"* unticked and keep your connections and values. The new **8 Amendment notice** flow arrives **Off**.
> 4. Run **0 Setup** again (4.3) and wait for *"ESG 1-to-1 setup: complete"*. It adds the new date columns and the **1-to-1 Amendments** list, and retires challenges (the technicians lose access to the 1-to-1 Challenges list; open challenges close as *Not agreed*; nothing is deleted).
> 5. Turn on flows **2, 3, 4, 5, 6 and 8**, in that order. Leave 1 and 7 off.
> 6. Upload the new frontend files with **Replace** and **Reset layout cache** (4.5). The top bar shows **v5.1.0**. (If the page goes up before step 4, it still works and shows a banner asking you to run Setup.)
> 7. ⚙ **Settings → Team**, for each person: **Appraisal period start** (the date their last appraisal was actually completed), **Next appraisal** (check it) and **Next wellbeing** (*Needs your attention* can offer the date from their last assessment – **Use it** – but only you decide). **Save settings.**
> 8. Turn on **Flow 1**.
> 9. *(Optional, separate approval.)* The one-time test-data reset (§7).
> 10. *(Separate approval.)* Delete *7 Challenge* from the solution once you're happy – an import never deletes a flow.

1. [ ] Go to **make.powerautomate.com**, signed in as yourself.
2. [ ] Top right: pick the environment IT named for the flows.
3. [ ] Left menu: **Solutions**.
   - If it asks to create a **Dataverse database**, or *Solutions* is missing, use §4.8 instead.
4. [ ] **Import solution → Browse →** `ESGOneToOne_1_5_0_1_unmanaged.zip` → **Next → Next**.
5. [ ] **Connections.** For each of the three, pick yours or click **+ New connection**, signed in as you:

   | Connection reference | Connector |
   |---|---|
   | ESG 1-to-1 SharePoint | SharePoint |
   | ESG 1-to-1 Outlook | Office 365 Outlook |
   | ESG 1-to-1 Office 365 Users | Office 365 Users |

6. [ ] **Environment variables:**

   | Variable | Value |
   |---|---|
   | ESG 1-to-1 Site URL | `https://<tenant>.sharepoint.com/sites/<site>`, the full site address with **no trailing slash** |
   | ESG 1-to-1 Team Leader email | your email address |

7. [ ] **Import.** Wait for *"imported successfully"* (1–2 minutes). If there's an error, copy the text from **Solutions → History**.

### 4.3 Run the setup flow (5 min)

1. [ ] **Solutions → ESG 1-to-1 → ESG 1-to-1 | 0 Setup (run once) → Run → Run flow.** If asked, confirm the connections (**Continue**).
2. [ ] Wait for **Succeeded** (2–4 minutes).
3. [ ] Your inbox: **"ESG 1-to-1 setup: complete"**.
   - If it says *"complete with items to check"*, fix exactly what the email lists, then run setup again. It is safe to re-run.

**What setup does:**

{SETUP_DOES}

### 4.4 Check the result (5 min)

- [ ] **Site contents** shows the eleven lists: Technicians, Appraisal Expectations, Monthly Reviews, Manager Reviews, **Wellbeing Assessments**, 1-to-1 Settings, **1-to-1 Feedback**, **1-to-1 Amendments**, **1-to-1 Actions**, **1-to-1 Action Updates**, **1-to-1 Action Notices** (plus the retired 1-to-1 Challenges on a site upgraded from 1.3).
- [ ] Each list → ⚙ → **List settings → Permissions for this list**:
  - Technicians, Appraisal Expectations, Monthly Reviews, Manager Reviews and **Wellbeing Assessments** (health information) show **only** *Owners – Full Control*;
  - **1-to-1 Settings** shows Owners **and** `1-to-1 Technicians` – Read;
  - **1-to-1 Feedback** shows Owners **and** `1-to-1 Technicians` – **1-to-1 Submit**; its Advanced settings → *Item-level permissions* are "Read items that were created by the user" and "Create items and edit items that were created by the user";
  - **1-to-1 Amendments** (and the retired **1-to-1 Challenges**, if present) show **only** *Owners – Full Control*; Amendments has versioning on (100);
  - **1-to-1 Actions**, **1-to-1 Action Updates** and **1-to-1 Action Notices** show **only** *Owners – Full Control* (flow 9 later shares each action and each person's update folder individually).
- [ ] **1-to-1 Settings** contains **Bands**, **Reminders** and **ActionsTrackerStart** (the date and time Setup 1.5 first ran – never edit it).
- [ ] Spot check: Monthly Reviews → List settings → click **Q9FollowUp**. The address ends `Field=Q9FollowUp`, and it has three choices.

### 4.5 Install the frontend (15 min)

1. [ ] **Site contents → Welcome Web Part - Custom Layouts → Upload → Files.** Upload `one-to-one.html` and `one-to-one.js` together. Choose **Replace** if asked.
2. [ ] Open the **OneToOne** page → **Edit**.
   - If the Welcome web part isn't already in a **full-width section**, add one and move the web part into it.
   - Web part **pencil** → **Layout = ESG 1-to-1**. This replaces the probe layout.
   - Greeting, time, weather and search: **off**; on the Content tab, News posts and Slides **off**.
   - If you upload a newer version later: **Reset layout cache** in the web part settings, then republish.
   - **Republish.**
3. [ ] **Site Pages → OneToOne → … → Make homepage.**
4. [ ] **Group permissions.** `1-to-1 Technicians` has **Read on exactly these, nothing more** (set in Stage 0; setup did the last one):
   - the `OneToOne` page;
   - the *Welcome Web Part - Custom Layouts* library;
   - the *Welcome Web Part - Errors* list;
   - the 1-to-1 Settings list.

### 4.6 Your first look

- [ ] **Version check.** If you replaced `one-to-one.html` / `one-to-one.js` after first choosing the layout: web part settings (pencil) → **Reset layout cache** → **Republish**.
- [ ] Open the site. You should see:
  - a dark **top menu bar**: Team overview, Monthly reviews, Expectations, Actions and Wellbeing on the left; **Team (0) ▾** and **⚙ Settings** on the right;
  - the month, deadline, ↻ Refresh and **v5.1.0** at the far right;
  - a compact **Feedback** pill at the bottom-right;
  - **Team overview** showing six tiles at 0, then *"No technicians yet"* with four numbered getting-started steps.
- [ ] Click through **Monthly reviews**, **Expectations**, **Actions**, **Wellbeing** and **⚙ Settings**. Expect no errors (empty pages are normal for now).
- [ ] **Expectations:** add one *Shared* expectation. This proves item creation and the Yes/No column work in your tenant.
  - Leave it Active if it's a real one; otherwise **Retire** it.
  - The *Individual* expectation (Person column) is tested in Phase B: the Technician dropdown is empty until someone is in the Technicians list.
- [ ] **Width and size:** the menu bar and content fill the page width, and the page ends where its content ends.
  - On a large monitor the layout scales itself up.
  - Ctrl + / Ctrl − still adjusts it per person.

> You can't see the *technician* side with your own account: the page always recognises you as Team Leader. Phase B uses Test A for that.

### 4.7 Turn on flows 1–6 and 8–11

- [ ] Only after the Setup summary email (4.3). **Solutions → ESG 1-to-1** → select each flow → **Turn on**, in this order:
  1. **2 Review submitted**
  2. **3 Manager decision**
  3. **4 Team members** – gives people you add in ⚙ Settings → Team access to the page
  4. **5 Wellbeing** – wellbeing assessments
  5. **6 Feedback** – emails you what people send with the 💬 Feedback bubble
  6. **8 Amendment notice** – emails a person when you amend their 1-to-1 (it arrives Off)
  7. **9 Action access** – shares each action with its person, and gives each person their update folder
  8. **10 Action update** – records the updates people send on their actions
  9. **11 Action notice** – sends the action emails (E19, E20)
  10. **1 Daily scheduler** – last, so its first run finds everything ready
- [ ] There is no Flow 7 from 1.4. If an old *7 Challenge* is still listed (after an upgrade), leave it **Off**.
- [ ] **0 Setup** stays On (it only runs when you click Run).

### 4.8 Fallback: no Dataverse (legacy packages)

Only if §4.2 step 3 showed that solutions can't be imported:

{LEGACY}

---

## 5. Phase B — Walk through it as a technician (test accounts only, 1–2 h)

- [ ] **⚙ Settings → Team → + Add a person:** search for **Test A** → **Add**, then the same for **Test B**. No SharePoint lists or groups to edit by hand.
  - Within about 5 minutes Flow 4 adds each to the `1-to-1 Technicians` group and emails you *"1-to-1: access given to …"* (test F32).
- [ ] **Reload the page.** The top bar shows **Team (2) ▾**, and opening it lists Test A and Test B.
- [ ] **⚙ Settings → Team:** give each a **Next appraisal** – type it (e.g. `15/07/2027`) or pick it with 📅 (use a summer date) – plus a **Next wellbeing** date and an **Appraisal period start**. Save, reopen, and check they **still show the same day** (tests F27, F31 and R16).
- [ ] **Expectations:** add one **Individual** expectation for Test A. This proves the Person column works; *Expectations per technician* shows Test A with one more than Test B.
- [ ] **Power Automate → Flow 1 → Run.** Test A and Test B each get *"Your 1-to-1 for … is ready"*.
  - **Run it again:** no new items, no second email.
- [ ] **As you:** Team overview lists Test A and Test B.
- [ ] **As Test A** (InPrivate window):
  - Open the email link. Their form opens. The top menu bar has only *This month · My actions · My appraisal · Wellbeing · Appraisal report*, with no Team menu or Settings.
  - Fill in a few answers. Within about 5 seconds the top shows **Saved hh:mm** – no need to click Save. Reload: the answers are kept (F43).
  - **Clearer forms (F45):** required questions not yet answered have a subtle orange border (Yes/No and rating buttons too). Press **Submit** with gaps: they turn red and the page jumps to the first; each clears as you answer it. Optional questions are never marked.
  - **Words, not numbers (F46):** the expectation rating buttons read *Not started · Early progress · On track · Strong progress · Fully met*, and the side panel shows no scores.
  - Complete everything → **Submit**. Within about 5 minutes the review becomes read-only.
- [ ] **As you:** *"1-to-1 submitted"* arrives. Monthly reviews shows a badge, and **Team ▾** shows a red dot next to Test A.
  - Open the review. Add feedback per question and a 1–5 rating per expectation.
  - Add **two actions**, each with owner, close-out criteria and deadline.
  - Click **Complete & release**.
- [ ] **As you:** Team ▾ → Test A → **Generate appraisal report** → **Print / Save as PDF**. A print window opens. If not, allow pop-ups for the SharePoint site.
- [ ] **As Test A:** *"Feedback … is ready"* arrives. Feedback, ratings and actions are visible; **My appraisal** and **Appraisal report** work.
- [ ] **Action tracking (F57–F61):**
  - **As you:** within about 5 minutes of the release, the **Actions** page lists both actions; the review says *In the action tracker ✓*.
  - **As Test A:** **My actions** shows them. On their own one, type an update → **Send update**. It says *received – being recorded*; after Refresh (within about 5 minutes) *recorded ✓*.
  - **As Test A:** **Ready to close… → Ask to close**. **As you:** *"Close-out requested"* arrives; *Needs your attention* lists it. Click **Not yet…** with a reason; Test A sees the reason; they ask again; you **Agree & close**.
  - **As you:** **Update…** on the other action → change the deadline: it won't save without a reason; with one, it shows the new date *(was <original>)*. Open **History** – every step is there with who and when.
  - Bands and words only – no numeric ratings or scores in the outlook, *My appraisal*, the charts or the report; the trend reads *Improving / Steady / Dipping / Too early to tell* (F46). Your own screens still show the numbers.
  - *To move up to <next band>, focus on:* three expectations (*Keep it up* in the top band), also in report section 1 *Areas to focus on* (F47).
- [ ] **Amendment (F23):**
  - **As you:** open Test A's released review. It is read-only. Click **Amend…**, change one feedback field → **Save amendment…**. It lists the changed field and won't save without a reason; give one.
  - **As Test A:** within about 5 minutes *"Your 1-to-1 for <month> has been amended"* arrives – the month and a link, nothing else – and the review shows the new value.
  - **As you:** the review shows **Amended (1)** (only you see it); no *"1-to-1 submitted"* email arrived.
- [ ] **Appraisal completed (F49):** ⚙ Settings → Team → **Mark appraisal completed…** for Test A (today, and a next due date a year on) → confirm. Test A's profile now has a period selector. Then **Undo last**.
- [ ] **Wellbeing (tests F37–F39):**
  - **As you:** **Wellbeing → + Start** for Test A, with a meeting date.
  - **As Test A:** *"Your wellbeing assessment is ready to complete"* arrives → open it → work through the steps → **Submit**.
  - **As you:** *"Wellbeing assessment submitted"* arrives → open it → add an arrangement and notes → **Complete** → confirm or set the **next wellbeing date** (it is never set automatically).
  - **As Test A:** *"Your wellbeing assessment record is ready"* arrives; the agreed record shows; **Print / Save as PDF** works.
- [ ] **Month two – "Since last time" and "Outstanding actions" (F44, F65):** with Flow 1's test date set to the next month (§6), Test A's new form starts with **Since last time** (what they raised and your replies) and **Outstanding actions** (open actions and what changed). Test A clicks **Progress satisfactory** on one; you see it on your review of that month.
- [ ] **Feedback (tests F41, P25):**
  - **As Test A:** click 💬 **Feedback** (bottom-right) → **Fault** → type a short message → **Send**. It says *"Thank you. Your feedback has been submitted to your Team Leader."*
  - **As you:** within about 5 minutes *"1-to-1 feedback (Fault): …"* arrives, naming Test A, the page and the version.
- [ ] **As Test B:**
  - nothing of Test A's is visible anywhere, including Wellbeing, **My actions** and Test A's feedback;
  - opening Test A's link (`?reviewId=` with A's number) shows **"Review not found"**, and `?wbId=` with A's assessment number shows **"Assessment not found"**.

---

## 6. Phase C — Security and functional test record (6–8 h, can span several days)

**Pass mark:**

- Every **P** test is mandatory: any failure **blocks go-live** until fixed (§10).
- F tests must pass too.

**Test date.** Tests F13–F20 change the date Flow 1 believes it is:

1. Open Flow 1 → **Edit** → **Initialize varTestDate** → set the value (e.g. `2026-11-20T07:00:00`).
2. **Save**, then **Run**.
3. **Clear the value afterwards and save.**

{H_INTRO}

### 6.1 Functional tests (F1–F68)

{F_TABLE}

### 6.2 Privacy and bypass tests (P1–P36) — all mandatory

{P_TABLE}

### 6.3 Questions for IT (record the answers)

| # | Question | Answer | Date |
|---|---|---|---|
| IT-a | Is *limited-access user permission lockdown mode* **off** for this site? (If it's on, technicians' page calls can fail.) |  |  |
| IT-b | IT-11: does the Welcome web part send any page or list data outside the tenant (licence checks, telemetry)? |  |  |

---

## 7. Phase D — Clean down (30 min)

- [ ] Delete every **test item** in Monthly Reviews, Manager Reviews, Wellbeing Assessments, 1-to-1 Feedback and 1-to-1 Amendments, and the test expectations.
- [ ] Empty the **site recycle bin**, then the **second-stage recycle bin**.
- [ ] **Lots of test data, or test data from earlier versions?** Use the one-time **reset tool** instead, only with written approval:
  1. Export the lists to CSV, and **turn Flow 1 off**.
  2. Upload `one-to-one-reset.html` and `one-to-one-reset.js` to *Welcome Web Part - Custom Layouts*, make a temporary Owners-only page with the Welcome web part and choose the layout **ESG 1-to-1 reset tool**.
  3. Tick the test records (filter by person or month; expectations only if you choose them) → **Download backup** → **Run the dry run** → tick *Flow 1 is turned off*, type the phrase shown (e.g. `DELETE 42 RECORDS`) → **Delete**.
  4. Check the result and **Download the report**. Deleted records go to the site recycle bin (restorable for about 93 days).
  5. Delete the temporary page and both reset files, then turn Flow 1 back on.

  It never touches the Technicians list (people, dates, period history), 1-to-1 Settings, permissions, groups or flows.
- [ ] **⚙ Settings → Team → Deactivate** Test A and Test B (Flow 4 removes them from the group), then delete their two items from the **Technicians** list. Team ▾ goes back to Team (0).
- [ ] Flow 1: **Initialize varTestDate** is **empty** (Save).

## 8. Phase E — Prepare and invite the real team (1–2 h)

- [ ] **⚙ Settings → Team → + Add a person** for each technician (Flow 4 gives them access; no invitation email is sent by SharePoint).
- [ ] **⚙ Settings → Team:** enter each person's **Next appraisal**, **Next wellbeing** and **Appraisal period start** (the date their last appraisal was actually completed) → Save. Optionally **⚙ Settings → Reminders / Names** to change the reminder days or the on-screen role words.
- [ ] **Expectations page:** enter the real **Shared** expectations and each person's **Individual** ones for **2026/27**. Do this after the Technicians list, so people appear in the Technician dropdown.
  - There is **no limit** on how many (a warning shows above 30 on one review). Shared always come first.
  - An expectation's year is the first appraisal year it applies to: it carries on into later years until you retire it. One entered early for next year stays hidden until that appraisal year starts (1 April).
- [ ] **Invite.** Send the welcome email below from your own mailbox, or brief the team in person.
  - It only tells them what's coming.
  - Access comes solely through each person's own monthly link, which works only for them.

### 8.1 Welcome email (copy, edit, send)

{WELCOME}

## 9. Phase F — Go live (15 min)

- [ ] **Flow 1 → Run** to start the current month now, or leave it to run automatically at **07:00 on the 1st**.
- [ ] Each technician has **exactly one** review and got **exactly one** email.
- [ ] **Team overview** lists everyone.

> ✔ **The system is fully implemented** when this phase is done and every P test in Phase C passed.

---

## 10. Phase G — First month, then the routine

| When | Check / do |
|---|---|
| Go-live day | Flow 1 run history is green; nobody is missing from Team overview |
| As submissions arrive | *"1-to-1 submitted"* email → open the review → feedback, ratings, actions → **Complete & release** |
| Actions (any time) | **Actions** page: the *What needs you* tiles – close-out requests (**Agree & close** / **Not yet…**), overdue, due soon, waiting, no recent progress, problems. Only you close or reopen. |
| E19 / E20 emails | Someone asked you to close an action (open it from the email), or a deadline is 5 working days away (you're copied) |
| Something needs changing after submission | Open the review → **Amend…** (released) or **Amend answers…** (submitted) → **Save amendment…** with a reason. The person gets a short email; *Needs your attention* shows any email that didn't go (**Resend**). |
| After each appraisal | ⚙ Settings → Team → **Mark appraisal completed…** (the actual date and the next due date) |
| E12 / E17 reminders | A wellbeing or appraisal date is 7 days away, due today or overdue – arrange it, or change the date in ⚙ Settings → Team |
| Your reminder days (default 20th and 27th) | Reminders went to anyone who hasn't submitted |
| Next 1st | New month is created; last month's unsubmitted reviews show **Overdue**; use **Close – not submitted** where needed |

### 10.1 Operating procedures

{PROCEDURES}

---

## 11. Troubleshooting

### 11.1 Import and setup

{IMP_TROUBLE}

### 11.2 Day to day

{CFG_TROUBLE}

### 11.3 Getting a fix

{FIX}

---

## 12. Appendix

### 12.1 Names you must never change

{NAMES}

### 12.2 What each flow does

| Flow | Starts | Job |
|---|---|---|
| ESG 1-to-1 \| 0 Setup (run once) | by hand | Creates or completes the lists, columns, indexes, views, permission levels, group, list permissions and the Bands and Reminders settings. Safe to re-run; never deletes data; emails a summary. |
| ESG 1-to-1 \| 1 Daily scheduler | 07:00 UK daily | Creates each month's reviews (with every expectation that applies, Shared first), makes them private, emails each technician their own link, sends reminders on your reminder days (default 20th and 27th), emails you 7 days before, on the day and weekly while overdue for each person's next wellbeing (E12) and appraisal (E17) dates, and reports problems to you – including a submitted review that was never processed. From 1.5 it also adds agreed actions Flow 3 couldn't, creates the E20 deadline reminders, and puts stuck action updates and failed action emails back in the queue |
| ESG 1-to-1 \| 2 Review submitted | when a technician submits | Locks the review (read-only), stamps the trusted record, emails you. Your own amendments are not treated as a new submission. |
| ESG 1-to-1 \| 3 Manager decision | when you Complete / Close | Validates, locks, releases your feedback to that technician only, emails them; from 1.5, adds the review's agreed actions to the action tracker (once each) |
| ESG 1-to-1 \| 4 Team members | when you add or deactivate someone (checks every 5 min) | Gives or removes access to the page; removes anyone in the group without an Active row; emails you |
| ESG 1-to-1 \| 5 Wellbeing | as a wellbeing assessment moves on (checks every 5 min) | Shares it with that person only, locks it when they submit, stamps completion, emails you both |
| ESG 1-to-1 \| 6 Feedback | when someone uses the 💬 Feedback bubble (checks every 5 min) | Makes the item Owners-only and emails it to you; if that fails, marks it Failed and it shows in *Needs your attention* |
| ESG 1-to-1 \| 8 Amendment notice | when you save an amendment, or click Resend (checks every 5 min) | Emails the person that their 1-to-1 for that month was amended (E16: the month and a link, never the change or the reason); marks it Sent, or Failed and tells you |
| ESG 1-to-1 \| 9 Action access | when an action is added or changed (checks every 5 min) | Makes each action readable by exactly its person, gives each person their own update folder, and removes any other access it finds |
| ESG 1-to-1 \| 10 Action update | when someone sends an update on an action (checks every 5 min) | Checks it really came from that person's own folder and follows the rules, records it on the action (never over your changes), locks it, and queues E19 for a close-out request |
| ESG 1-to-1 \| 11 Action notice | when an action email is waiting (checks every 5 min) | Sends E19 (close-out requested, to you) and E20 (deadline in 5 working days, to the person, copied to you); records Sent or Failed |

There is no Flow 7 from 1.4 (the challenge flow was removed).

### 12.3 Accepted risks (from the implementation plan)

{RISKS}
'''

md = GUIDE
for key, val in {'SETUP_DOES': setup_does, 'LEGACY': demote(legacy, 2), 'H_INTRO': H_intro,
                 'F_TABLE': add_result_columns(F_table), 'P_TABLE': add_result_columns(P_table),
                 'WELCOME': welcome, 'PROCEDURES': procedures, 'IMP_TROUBLE': imp_trouble,
                 'CFG_TROUBLE': cfg_trouble, 'FIX': fix_prompt, 'NAMES': names,
                 'RISKS': '| # | Risk | Mitigation |\n|---|---|---|\n' + '\n'.join(l for l in risks.split('\n')
                                                                               if not l.startswith('| #') and not l.startswith('|---'))}.items():
    md = md.replace('{' + key + '}', val)
md = delink(md)
md = md.replace(' 🔒', '').replace('🔒', '')
# cross-references copied from the source files → this guide's own numbering
for a, b in (('## Phase 0 — Get and check the files', '## 3. Phase 0 — Get and check the files'),
             ("Use this only if §1 showed that solutions can't be imported.\n\n", ''),
             ('Continue at §3 (run flow 0), then §4 and §5.', 'Continue at §4.3 (run the setup flow), then §4.4–§4.7.'),
             ('blocks go-live** until fixed (§10)', 'blocks go-live** until fixed (§11.3)'),
             ('matching `../01-sharepoint/lists.md` exactly', 'matching the agreed list specification exactly'),
             ('Enter both values in step 2.3', 'Enter both values in §4.2 step 6'),
             ('Every internal name in `01-sharepoint/lists.md`', 'Every column name the setup flow created (List settings shows them)')):
    assert a in md, 'cross-reference not found: ' + a
    md = md.replace(a, b)
assert '{' + 'F_TABLE' + '}' not in md and not re.search(r'\{[A-Z_]{4,}\}', md), 'unfilled placeholder'

out_md = os.path.join(ROOT, 'FINALISATION_GUIDE.md')
with open(out_md, 'w', encoding='utf-8') as f:
    f.write(md)

# ---- Word version
docx_md = md.replace('- [ ] ', '- ☐ ')
docx_md = re.sub(r'^(\s*)(\d+)\. \[ \] ', r'\1\2. ☐ ', docx_md, flags=re.M)
tmp = os.path.join(ROOT, '.guide-tmp.md')
with open(tmp, 'w', encoding='utf-8') as f:
    f.write(docx_md)
out_docx = os.path.join(ROOT, 'ESG-1to1-Finalisation-Guide.docx')
subprocess.run(['pandoc', '-f', 'gfm', '-t', 'docx', '-o', out_docx, tmp], check=True)
os.remove(tmp)

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402,F401
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

doc = Document(out_docx)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.6)
base = doc.styles['Normal'] if 'Normal' in [st.name for st in doc.styles] else None
for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    try:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(10.5)
    except KeyError:
        pass
for name, size in (('Title', 22), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)):
    try:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x13, 0x51, 0x6B)
    except KeyError:
        pass


def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'A9B4BD')
        b.append(e)
    tblPr.append(b)


TEXT_W = 21.0 - 2 * 1.8   # cm
WIDTHS = {   # header text of the first two columns -> column widths (cm); other tables share the width by content
    ('#', 'Test'): [1.3, 5.5, 5.4, 1.6, 1.5, 2.1],
    ('#', 'Attempt by Test A'): [1.2, 2.6, 5.1, 4.1, 1.5, 1.4, 1.5],
}


def set_widths(tbl):
    from docx.shared import Cm as _Cm
    hdr = [c.text.strip() for c in tbl.rows[0].cells]
    widths = WIDTHS.get(tuple(hdr[:2]))
    if not widths or len(widths) != len(hdr):
        lens = [max(7, min(60, max(len(r.cells[i].text) for r in tbl.rows))) for i in range(len(hdr))]
        widths = [TEXT_W * l / sum(lens) for l in lens]
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(int(TEXT_W / 2.54 * 1440)))
    tw.set(qn('w:type'), 'dxa')
    tblPr.append(tw)
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w / 2.54 * 1440)))
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = _Cm(w)


for tbl in doc.tables:
    borders(tbl)
    set_widths(tbl)
    for i, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'E4EDF2')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
doc.core_properties.title = 'ESG Monthly Technician 1-to-1 — Finalisation Guide'
doc.core_properties.author = 'ESG 1-to-1'
doc.save(out_docx)
print('written', os.path.relpath(out_md, ROOT), 'and', os.path.relpath(out_docx, ROOT))
