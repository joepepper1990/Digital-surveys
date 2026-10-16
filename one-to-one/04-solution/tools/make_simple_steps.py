"""Builds the printable Word version of SIMPLE_UPGRADE_STEPS.md (plain-language upgrade steps).

    python3 04-solution/tools/make_simple_steps.py        (from anywhere)

Output: one-to-one/ESG-1to1-Simple-Upgrade-Steps.docx. Edit the Markdown, then re-run this.
"""
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))   # one-to-one/
SRC = os.path.join(ROOT, 'SIMPLE_UPGRADE_STEPS.md')
OUT = os.path.join(ROOT, 'ESG-1to1-Simple-Upgrade-Steps.docx')


def add_runs(p, text):
    """**bold** and *italic*; everything else as plain text."""
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part:
            continue
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*'):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


def build():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(12)
    for ln in open(SRC, encoding='utf-8').read().split('\n'):
        if not ln.strip():
            continue
        indent = len(ln) - len(ln.lstrip(' '))
        s = ln.strip()
        if s.startswith('# '):
            doc.add_heading(s[2:], level=0)
        elif s.startswith('## '):
            h = doc.add_heading(s[3:], level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x13, 0x51, 0x6B)
        elif s.startswith('- [ ] '):
            p = doc.add_paragraph()
            p.add_run('☐  ')
            add_runs(p, s[6:])
        elif re.match(r'^\d+\. ', s):
            n, rest = s.split('. ', 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18 if indent == 0 else 36)
            p.paragraph_format.first_line_indent = Pt(-18)
            p.add_run('☐ ' + n + '.  ').bold = True
            add_runs(p, rest)
        elif s.startswith('- '):
            p = doc.add_paragraph(style='List Bullet' if indent == 0 else 'List Bullet 2')
            add_runs(p, s[2:])
        else:
            add_runs(doc.add_paragraph(), s)
    doc.save(OUT)
    return OUT


if __name__ == '__main__':
    print('written', build())
